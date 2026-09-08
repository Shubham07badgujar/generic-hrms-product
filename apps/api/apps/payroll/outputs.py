"""
Payroll documents: payslip PDF, salary register, bank advice.

All three are generated from the STORED payslip rather than recomputed. A
payslip PDF that recalculated its own figures could disagree with the payslip
it claims to represent the moment a component's configuration changed, and the
employee holding the PDF would be right while the system was wrong.

Bank account numbers are encrypted at rest and masked in every API response.
The NEFT file is the one place a full account number legitimately appears —
without it the bank cannot pay anyone — so it is export-permission gated and
audited at the point of download.
"""

from __future__ import annotations

import io
from decimal import Decimal

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

from .models import ComponentType, StatutoryKind

ZERO = Decimal("0.00")

MONTHS = [
    "", "January", "February", "March", "April", "May", "June",
    "July", "August", "September", "October", "November", "December",
]


def _organisation_name() -> str:
    from apps.organization.models import OrgSettings

    settings_row = OrgSettings.objects.first()
    return settings_row.name if settings_row else "Organisation"


# ---------------------------------------------------------------------------
# Payslip PDF
# ---------------------------------------------------------------------------


#: The letterhead palette the offer letters established.
BRAND = colors.HexColor("#232B7C")
GREEN = colors.HexColor("#A6CE39")
TINT = colors.HexColor("#f3f8e4")
GRID = colors.HexColor("#c9ced6")
INK = colors.HexColor("#1a1a1a")

_UNITS = ["", "One", "Two", "Three", "Four", "Five", "Six", "Seven", "Eight", "Nine",
          "Ten", "Eleven", "Twelve", "Thirteen", "Fourteen", "Fifteen", "Sixteen",
          "Seventeen", "Eighteen", "Nineteen"]
_TENS = ["", "", "Twenty", "Thirty", "Forty", "Fifty", "Sixty", "Seventy", "Eighty", "Ninety"]


def _two(n: int) -> str:
    return _UNITS[n] if n < 20 else (_TENS[n // 10] + (f" {_UNITS[n % 10]}" if n % 10 else ""))


def _in_words(amount: Decimal) -> str:
    """Indian-system words: Rupees Twelve Thousand Six Hundred Eighty Eight only."""
    rupees = int(amount)
    paise = int(round((amount - rupees) * 100))
    if rupees == 0:
        words = "Zero"
    else:
        parts = []
        for value, label in ((rupees // 10**7, "Crore"), (rupees // 10**5 % 100, "Lakh"),
                             (rupees // 1000 % 100, "Thousand"), (rupees // 100 % 10, "Hundred")):
            if value:
                parts.append(f"{_two(value)} {label}")
        if rupees % 100:
            parts.append(_two(rupees % 100))
        words = " ".join(parts)
    return f"Rupees {words}" + (f" and Paise {_two(paise)}" if paise else "") + " only"


def payslip_pdf(payslip) -> bytes:
    """
    A payslip the employee can read, check and file — the professional grid
    layout HR asked for: letterhead with the logo, a boxed details panel,
    earnings shown Standard vs Earned (so proration is visible, not
    mysterious), deductions alongside, and the net highlighted.

    Employer contributions stay a separate block, and only when the employee
    is actually enrolled in something — they were never in the person's pay.
    """
    from apps.recruitment.letters import _org, _scaled_image

    buffer = io.BytesIO()
    document = SimpleDocTemplate(
        buffer, pagesize=A4,
        leftMargin=14 * mm, rightMargin=14 * mm,
        topMargin=14 * mm, bottomMargin=14 * mm,
        title=f"Payslip {payslip.employee.employee_code}",
    )

    styles = getSampleStyleSheet()
    org_line = ParagraphStyle("org", parent=styles["Heading1"], fontSize=15,
                              textColor=BRAND, spaceAfter=0)
    sub = ParagraphStyle("s", parent=styles["Normal"], fontSize=9, textColor=colors.grey)
    note = ParagraphStyle("n", parent=styles["Normal"], fontSize=8, textColor=colors.grey,
                          alignment=1, spaceBefore=10)

    run = payslip.payroll_run
    employee = payslip.employee
    organisation = _org()
    org_name = organisation.name if organisation else "Organisation"
    legal_name = organisation.legal_name if organisation and organisation.legal_name else org_name
    period = f"{MONTHS[run.period_month]} {run.period_year}"
    if run.run_type != "regular":
        period += f" — {run.get_run_type_display()}"

    # ---- letterhead: logo left, names right, title band under a green rule
    logo = _scaled_image(organisation.logo if organisation else None,
                         max_width=52 * mm, max_height=17 * mm)
    head_right = [Paragraph(legal_name, org_line)]
    if legal_name != org_name:
        head_right.append(Paragraph(org_name, sub))
    header = Table(
        [[logo or "", head_right]],
        colWidths=[58 * mm, None],
        style=TableStyle([
            ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
            ("LEFTPADDING", (0, 0), (0, 0), 0),
            ("RIGHTPADDING", (-1, -1), (-1, -1), 0),
        ]),
    )
    title_band = Table(
        [[f"SALARY SLIP — {period.upper()}"]], colWidths=[182 * mm],
        style=TableStyle([
            ("BACKGROUND", (0, 0), (-1, -1), BRAND),
            ("TEXTCOLOR", (0, 0), (-1, -1), colors.white),
            ("FONTNAME", (0, 0), (-1, -1), "Helvetica-Bold"),
            ("FONTSIZE", (0, 0), (-1, -1), 10.5),
            ("ALIGN", (0, 0), (-1, -1), "CENTER"),
            ("TOPPADDING", (0, 0), (-1, -1), 5),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
            ("LINEBELOW", (0, 0), (-1, -1), 1.6, GREEN),
        ]),
    )

    story = [
        header,
        Spacer(1, 6),
        title_band,
        Spacer(1, 8),
        _details_grid(payslip, employee),
        Spacer(1, 8),
        _earnings_and_deductions(payslip),
        Spacer(1, 8),
        _net_band(payslip),
        # Only when there is one: an employee outside PF/ESIC/gratuity (the
        # structure decides) gets no employer-contribution block at all,
        # rather than a table of zeros implying benefits they do not have.
        *(
            [Spacer(1, 8), _employer_block(employer_lines, payslip.employer_contributions)]
            if (employer_lines := [
                (line.label, line.amount)
                for line in payslip.lines.all() if line.is_employer_side
            ])
            else []
        ),
        Paragraph(
            "This is a system-generated payslip and does not require a signature. "
            "Statutory deductions follow the rates in force for the period shown. "
            "Queries should go to the payroll team.",
            note,
        ),
    ]

    document.build(story)
    return buffer.getvalue()


def _num(value) -> str:
    """'21.00' -> '21', '0.5' -> '0.5', None -> em-dash (not tracked)."""
    if value is None or value == "":
        return "—"
    return f"{Decimal(str(value)).normalize():f}"


def _details_grid(payslip, employee) -> Table:
    """The boxed employee/period panel, two label-value pairs per row."""
    import calendar as _calendar
    from itertools import zip_longest

    run = payslip.payroll_run
    month_days = _calendar.monthrange(run.period_year, run.period_month)[1]
    #: Frozen at process time. An older payslip without it still prints the
    #: figures the payslip itself carries; the rest show as "not tracked".
    summary = payslip.attendance_summary or {}
    left = [
        ("Employee Name", employee.full_name),
        ("Employee Code", employee.employee_code),
        ("Designation", employee.designation.title if employee.designation else "—"),
        ("Department", employee.department.name if employee.department else "—"),
        # Masked, never the full numbers — a payslip is emailed and printed.
        ("PAN", employee.pan_masked or "—"),
        ("Bank A/C", employee.bank_account_masked or "—"),
        ("Bank Name", employee.bank_name or "—"),
        ("UAN", employee.uan or "—"),
        ("ESIC No.", employee.esic_number or "—"),
    ]
    right = [
        ("Salary for Month", f"{MONTHS[run.period_month]} {run.period_year}"),
        ("Center / Location", payslip.location.name if payslip.location else "—"),
        ("Days in Month", str(summary.get("days_in_month", month_days))),
        ("Working Days", _num(summary.get("working_days"))),
        ("Weekly Offs (paid)", _num(summary.get("weekly_offs"))),
        ("Days Attended", _num(summary.get("days_attended"))),
        ("Absent Days", _num(summary.get("absent_days"))),
        ("Half Days", _num(summary.get("half_days"))),
        ("Paid Leave Taken", _num(summary.get("paid_leave_taken"))),
        ("Leave Balance", _num(summary.get("leave_balance"))),
        ("Loss of Pay (LWP)", f"{payslip.lop_days:g}"),
        ("Paid Days", f"{payslip.paid_days:g}"),
    ]
    rows = [
        [l1, v1, l2, v2]
        for (l1, v1), (l2, v2) in zip_longest(left, right, fillvalue=("", ""))
    ]
    table = Table(rows, colWidths=[30 * mm, 61 * mm, 32 * mm, 59 * mm])
    table.setStyle(TableStyle([
        ("GRID", (0, 0), (-1, -1), 0.5, GRID),
        ("BOX", (0, 0), (-1, -1), 1, BRAND),
        ("FONTSIZE", (0, 0), (-1, -1), 8.5),
        ("BACKGROUND", (0, 0), (0, -1), TINT),
        ("BACKGROUND", (2, 0), (2, -1), TINT),
        ("FONTNAME", (0, 0), (0, -1), "Helvetica-Bold"),
        ("FONTNAME", (2, 0), (2, -1), "Helvetica-Bold"),
        ("TEXTCOLOR", (0, 0), (-1, -1), INK),
        ("TOPPADDING", (0, 0), (-1, -1), 3.5),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 3.5),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
    ]))
    return table


def _earnings_and_deductions(payslip) -> Table:
    """
    Earnings (Standard | Earned) beside Deductions, one boxed table.

    Standard is the structure's full monthly figure; Earned is what the paid
    days produced — the proration is visible instead of mysterious. A line
    with no structure counterpart (a bonus, an arrear) shows "—" as standard.
    """
    from apps.payroll.models import SalaryComponent

    standard_by_component = {}
    if payslip.salary_structure_id:
        standard_by_component = {
            line.component_id: line.monthly_amount
            for line in payslip.salary_structure.lines.all()
        }
    earned_by_component = {
        line.component_id: line.amount
        for line in payslip.lines.all()
        if line.component_id and not line.is_employer_side
        and line.component_type in (ComponentType.EARNING, ComponentType.REIMBURSEMENT)
    }

    # Every earning component in the catalogue appears, zeroes included —
    # HR wants the slip to show the full standard shape of pay, not only the
    # parts this person happens to receive.
    earning_rows: list[list[str]] = []
    for component in SalaryComponent.objects.filter(
        is_active=True, component_type__in=(ComponentType.EARNING, ComponentType.REIMBURSEMENT),
    ).order_by("display_order", "code"):
        standard = standard_by_component.get(component.pk, ZERO)
        earned = earned_by_component.get(component.pk, ZERO)
        earning_rows.append([component.name, f"{standard:,.2f}", f"{earned:,.2f}"])
    # Lines with no catalogue component — a bonus, an arrear — after them.
    for line in payslip.lines.all():
        if (line.component_id is None and not line.is_employer_side
                and line.component_type in (ComponentType.EARNING, ComponentType.REIMBURSEMENT)):
            earning_rows.append([line.label, "—", f"{line.amount:,.2f}"])

    # Deductions: the four statutory heads always print, zeroes included.
    deduction_lines = {
        line.label: line.amount
        for line in payslip.lines.all()
        if not line.is_employer_side
        and line.component_type in (ComponentType.DEDUCTION, ComponentType.STATUTORY_DEDUCTION)
    }
    deduction_rows: list[list[str]] = []
    for label in ("Provident Fund (employee)", "ESI (employee)", "Professional Tax", "TDS"):
        deduction_rows.append([label, f"{deduction_lines.pop(label, ZERO):,.2f}"])
    for label, amount in deduction_lines.items():  # loans, other deductions
        deduction_rows.append([label, f"{amount:,.2f}"])

    body_rows = max(len(earning_rows), len(deduction_rows), 1)
    data = [["Earnings (INR)", "Standard", "Earned", "Deductions (INR)", "Amount"]]
    for index in range(body_rows):
        earning_cells = earning_rows[index] if index < len(earning_rows) else ["", "", ""]
        deduction_cells = deduction_rows[index] if index < len(deduction_rows) else ["", ""]
        data.append(earning_cells + deduction_cells)
    data.append([
        "Total", "", f"{payslip.gross_earnings:,.2f}",
        "Total", f"{payslip.total_deductions:,.2f}",
    ])

    table = Table(data, colWidths=[44 * mm, 24 * mm, 26 * mm, 56 * mm, 32 * mm])
    table.setStyle(TableStyle([
        ("GRID", (0, 0), (-1, -1), 0.5, GRID),
        ("BOX", (0, 0), (-1, -1), 1, BRAND),
        ("LINEAFTER", (2, 0), (2, -1), 1, BRAND),  # the earnings/deductions split
        ("BACKGROUND", (0, 0), (-1, 0), BRAND),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("FONTSIZE", (0, 0), (-1, -1), 8.5),
        ("ALIGN", (1, 0), (2, -1), "RIGHT"),
        ("ALIGN", (4, 0), (4, -1), "RIGHT"),
        ("BACKGROUND", (0, -1), (-1, -1), TINT),
        ("FONTNAME", (0, -1), (-1, -1), "Helvetica-Bold"),
        ("TOPPADDING", (0, 0), (-1, -1), 3.5),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 3.5),
    ]))
    return table


def _net_band(payslip) -> Table:
    table = Table(
        [[f"Net Salary:  {payslip.net_pay:,.2f}", _in_words(payslip.net_pay)]],
        colWidths=[68 * mm, 114 * mm],
    )
    table.setStyle(TableStyle([
        ("BOX", (0, 0), (-1, -1), 1, BRAND),
        ("BACKGROUND", (0, 0), (0, 0), GREEN),
        ("FONTNAME", (0, 0), (0, 0), "Helvetica-Bold"),
        ("FONTSIZE", (0, 0), (0, 0), 11),
        ("TEXTCOLOR", (0, 0), (0, 0), BRAND),
        ("FONTSIZE", (1, 0), (1, 0), 8.5),
        ("TEXTCOLOR", (1, 0), (1, 0), INK),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("TOPPADDING", (0, 0), (-1, -1), 6),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 6),
    ]))
    return table


def _employer_block(rows: list[tuple[str, Decimal]], total: Decimal) -> Table:
    data = [["Employer Contributions (not deducted from your pay)", "Amount (INR)"]]
    data += [[label, f"{amount:,.2f}"] for label, amount in rows]
    data.append(["Total Employer Cost", f"{total:,.2f}"])
    table = Table(data, colWidths=[142 * mm, 40 * mm])
    table.setStyle(TableStyle([
        ("GRID", (0, 0), (-1, -1), 0.5, GRID),
        ("BOX", (0, 0), (-1, -1), 1, BRAND),
        ("BACKGROUND", (0, 0), (-1, 0), TINT),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("FONTNAME", (0, -1), (-1, -1), "Helvetica-Bold"),
        ("FONTSIZE", (0, 0), (-1, -1), 8.5),
        ("ALIGN", (1, 0), (1, -1), "RIGHT"),
        ("TOPPADDING", (0, 0), (-1, -1), 3.5),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 3.5),
    ]))
    return table


# ---------------------------------------------------------------------------
# Salary register
# ---------------------------------------------------------------------------


REGISTER_COLUMNS = [
    "Employee code", "Name", "Department", "Location", "State",
    "Paid days", "LOP days", "Gross earnings",
    "PF (employee)", "ESI (employee)", "Professional Tax", "TDS",
    "Total deductions", "Net pay",
    "PF (employer)", "ESI (employer)", "Gratuity provision", "Employer cost",
]


def salary_register(run) -> tuple[bytes, str]:
    """One row per employee — the sheet finance reconciles the run against."""
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = f"{run.period_year}-{run.period_month:02d}"

    sheet.append(REGISTER_COLUMNS)
    for cell in sheet[1]:
        cell.font = Font(bold=True)
        cell.alignment = Alignment(horizontal="center")

    payslips = (
        run.payslips.select_related("employee__department", "location")
        .prefetch_related("statutory_contributions")
        .order_by("employee__employee_code")
    )

    for payslip in payslips:
        contributions = {c.kind: c for c in payslip.statutory_contributions.all()}

        def employee_side(kind) -> Decimal:
            row = contributions.get(kind)
            return row.employee_amount if row else ZERO

        def employer_side(kind) -> Decimal:
            row = contributions.get(kind)
            return row.employer_amount if row else ZERO

        sheet.append([
            payslip.employee.employee_code,
            payslip.employee.full_name,
            payslip.employee.department.name if payslip.employee.department else "",
            payslip.location.name if payslip.location else "",
            payslip.state,
            float(payslip.paid_days), float(payslip.lop_days),
            float(payslip.gross_earnings),
            float(employee_side(StatutoryKind.PF)),
            float(employee_side(StatutoryKind.ESI)),
            float(employee_side(StatutoryKind.PT)),
            float(employee_side(StatutoryKind.TDS)),
            float(payslip.total_deductions),
            float(payslip.net_pay),
            float(employer_side(StatutoryKind.PF)),
            float(employer_side(StatutoryKind.ESI)),
            float(employer_side(StatutoryKind.GRATUITY)),
            float(payslip.employer_contributions),
        ])

    # A totals row, because the first thing anyone does with this sheet is add
    # up the net pay column and compare it to the bank file.
    last = sheet.max_row
    if last > 1:
        sheet.append([])
        totals = ["", "TOTAL", "", "", "", "", ""] + [
            f"=SUM({get_column_letter(col)}2:{get_column_letter(col)}{last})"
            for col in range(8, len(REGISTER_COLUMNS) + 1)
        ]
        sheet.append(totals)
        for cell in sheet[sheet.max_row]:
            cell.font = Font(bold=True)

    for index, heading in enumerate(REGISTER_COLUMNS, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = max(12, len(heading) + 2)
    sheet.freeze_panes = "C2"

    buffer = io.BytesIO()
    workbook.save(buffer)
    filename = f"salary-register-{run.period_year}-{run.period_month:02d}.xlsx"
    return buffer.getvalue(), filename


# ---------------------------------------------------------------------------
# Bank advice
# ---------------------------------------------------------------------------


NEFT_COLUMNS = ["EMPLOYEE_CODE", "NAME", "BANK_ACCOUNT", "IFSC", "AMOUNT", "REMARK"]


def neft_advice(run) -> tuple[str, Decimal]:
    """
    Tab-separated payment instruction for the bank.

    Employees with no net pay are omitted rather than sent as zero rows: a bank
    file is a set of instructions to move money, and a zero instruction is at
    best noise and at worst a rejected batch.

    Missing bank details are included as an explicit row with a blank account
    and a remark, so the payment run surfaces them instead of the employee
    silently going unpaid.
    """
    lines = ["\t".join(NEFT_COLUMNS)]
    total = ZERO

    payslips = run.payslips.select_related("employee").order_by("employee__employee_code")
    for payslip in payslips:
        if payslip.net_pay <= ZERO:
            continue
        employee = payslip.employee
        account = employee.bank_account_number or ""
        remark = f"SAL {run.period_year}-{run.period_month:02d}"
        if not account or not employee.bank_ifsc:
            remark = "MISSING BANK DETAILS — NOT PAYABLE"
        else:
            total += payslip.net_pay

        lines.append("\t".join([
            employee.employee_code,
            employee.full_name,
            account,
            employee.bank_ifsc or "",
            f"{payslip.net_pay:.2f}",
            remark,
        ]))

    return "\n".join(lines) + "\n", total
